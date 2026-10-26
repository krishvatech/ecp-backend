"""
Bulk attendee import: superusers pre-register people for an event from a CSV/XLSX file.

Flow:
- preview: parse + validate the file and plan each row (no writes). The plan is
  cached under a plan_id for PLAN_TTL seconds.
- apply: a Celery job runs the plan. Every row is re-planned at apply time (state
  may have changed since the preview) and runs in its own transaction.

Per row:
- Account: an existing account (matched by email) is used as-is. A missing one is
  created in Cognito (email verified, so no verification code) and Django, and
  emailed a temporary password.
- Open event: register (reactivating a cancelled registration) via the invite helper.
- Apply event: reuse or create the person's application on the selected track and
  accept it via accept_track_application (roles, origins, tier, forms), with the
  acceptance email turned off.
- Both: add the type badge label; "Paid Attendee" is marked paid offline; send the
  registration confirmation email. Besides the built-in types, any active badge
  label of the event is accepted as a Type (label only).
"""
import csv
import io
import logging
import uuid

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from events.models import (
    EventApplication,
    EventApplicationTrackApplication,
    EventRegistration,
)

logger = logging.getLogger(__name__)
User = get_user_model()

MAX_ROWS = 1000
MAX_FILE_BYTES = 5 * 1024 * 1024
PLAN_TTL = 30 * 60
JOB_TTL = 24 * 60 * 60
PLAN_CACHE_KEY = "attendee_import:plan:{}"
JOB_CACHE_KEY = "attendee_import:job:{}"

TYPE_ATTENDEE = "attendee"
TYPE_SPONSOR = "sponsor"
TYPE_PAID_ATTENDEE = "paid_attendee"
TYPE_SPONSORED_GUEST = "sponsored_guest"
TYPE_LABELS = {
    TYPE_ATTENDEE: "Attendee",
    TYPE_SPONSOR: "Sponsor",
    TYPE_PAID_ATTENDEE: "Paid Attendee",
    TYPE_SPONSORED_GUEST: "Sponsored Guest",
}
TYPE_LABEL_COLORS = {
    TYPE_ATTENDEE: "#6366f1",
    TYPE_SPONSOR: "#f59e0b",
    TYPE_PAID_ATTENDEE: "#10b981",
    TYPE_SPONSORED_GUEST: "#ec4899",
}
# Any active badge label of the event is also accepted as a Type (label only)
TYPE_EVENT_LABEL = "event_label"
EVENT_LABEL_COLOR = "#9CA3AF"

REQUIRED_COLUMNS = ("email", "first_name", "last_name", "type")
OPTIONAL_COLUMNS = ("job_title", "company", "country", "payment_reference")
COLUMN_ALIASES = {
    "email_address": "email",
    "firstname": "first_name",
    "lastname": "last_name",
    "surname": "last_name",
    "attendee_type": "type",
    "registration_type": "type",
    "title": "job_title",
    "company_name": "company",
    "organisation": "company",
    "organization": "company",
    "location": "country",
    "country_region": "country",
}
FIELD_MAX_LENGTHS = {"first_name": 150, "last_name": 150, "job_title": 200, "company": 200, "country": 255, "payment_reference": 255}

TEMPLATE_HEADER = list(REQUIRED_COLUMNS + OPTIONAL_COLUMNS)
TEMPLATE_EXAMPLE = ["jane.doe@example.com", "Jane", "Doe", "Attendee", "Head of M&A", "Example Corp", "Singapore", ""]

# Row actions
ACTION_REGISTER = "register"
ACTION_ACCEPT = "accept_application"
ACTION_SKIP = "skip"
ACTION_ERROR = "error"

# Track application statuses that an import may accept
ACCEPTABLE_TRACK_STATUSES = {
    EventApplicationTrackApplication.STATUS_PENDING,
    EventApplicationTrackApplication.STATUS_PRE_APPROVED,
    EventApplicationTrackApplication.STATUS_WAITLISTED,
    EventApplicationTrackApplication.STATUS_DECLINED,
    EventApplicationTrackApplication.STATUS_CANCELLED,
}


class ImportFileError(ValueError):
    """The file (or the import as a whole) cannot be processed."""

    def __init__(self, message, code="invalid_file"):
        super().__init__(message)
        self.code = code


# ---------------------------------------------------------------------------
# Parsing + validation
# ---------------------------------------------------------------------------

def _normalize_header(value):
    key = str(value or "").strip().lower().replace("-", " ").replace("/", " ")
    key = "_".join(key.split())
    return COLUMN_ALIASES.get(key, key)


def _normalize_type(value):
    key = "_".join(str(value or "").strip().lower().replace("-", " ").split())
    return key if key in TYPE_LABELS else None


def _label_key(name):
    return " ".join(str(name or "").split()).lower()


def event_type_labels(event):
    """The event's active badge label names, usable as a Type besides the built-in types."""
    from events.models import EventBadgeLabel

    return list(
        EventBadgeLabel.objects.filter(event=event, is_active=True).order_by("name").values_list("name", flat=True)
    )


def resolve_event_types(rows, event):
    """Match a Type that is not built in against the event's active badge labels."""
    labels = {_label_key(name): name for name in event_type_labels(event)}
    for row in rows:
        if row["type"] or not row["type_raw"]:
            continue
        name = labels.get(_label_key(row["type_raw"]))
        if name:
            row["type"] = TYPE_EVENT_LABEL
            row["type_label"] = name
    return rows


def type_options(event):
    """Every Type value the event accepts: built-in types, then its own badge labels."""
    options = list(TYPE_LABELS.values())
    known = {_label_key(name) for name in options}
    options += [name for name in event_type_labels(event) if _label_key(name) not in known]
    return options


def _read_table(upload):
    name = (getattr(upload, "name", "") or "").lower()
    data = upload.read()
    if len(data) > MAX_FILE_BYTES:
        raise ImportFileError("File is larger than 5 MB.", code="file_too_large")

    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ImportFileError("Could not read the Excel file.")
        sheet = workbook.worksheets[0]
        table = [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        workbook.close()
        return table

    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        return [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text))]

    raise ImportFileError("Upload a .csv or .xlsx file.", code="unsupported_file_type")


def parse_file(upload):
    """
    Return the file's data rows as dicts:
    {"row": <spreadsheet row number>, "email", "first_name", ..., "errors": [...]}.
    Raises ImportFileError for a file that cannot be processed at all.
    """
    table = _read_table(upload)
    # Drop fully empty lines (common at the end of exported sheets)
    numbered = [(index + 1, cells) for index, cells in enumerate(table) if any(cells)]
    if not numbered:
        raise ImportFileError("The file is empty.", code="empty_file")

    _, header_cells = numbered[0]
    headers = [_normalize_header(cell) for cell in header_cells]
    missing = [column for column in REQUIRED_COLUMNS if column not in headers]
    if missing:
        raise ImportFileError(
            f"Missing required column(s): {', '.join(missing)}.", code="missing_columns"
        )

    data_rows = numbered[1:]
    if not data_rows:
        raise ImportFileError("The file has no attendee rows.", code="empty_file")
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(f"The file has more than {MAX_ROWS} rows.", code="too_many_rows")

    rows = []
    for row_number, cells in data_rows:
        values = {}
        for position, header in enumerate(headers):
            if header in REQUIRED_COLUMNS or header in OPTIONAL_COLUMNS:
                if header not in values or not values[header]:
                    values[header] = cells[position] if position < len(cells) else ""
        row = {"row": row_number}
        for column in REQUIRED_COLUMNS + OPTIONAL_COLUMNS:
            value = (values.get(column) or "").strip()
            limit = FIELD_MAX_LENGTHS.get(column)
            row[column] = value[:limit] if limit else value
        row["email"] = row["email"].lower()
        row["type_raw"] = row.pop("type")
        row["type"] = _normalize_type(row["type_raw"])
        row["type_label"] = TYPE_LABELS.get(row["type"], "")
        rows.append(row)
    return rows


def validate_rows(rows):
    """
    Add per-row errors (missing/invalid fields, duplicates) and warnings.

    A disposable email domain is only a warning: like admin Create User, staff
    provisioning may use any valid address (public signup still blocks them).
    """
    from users.disposable_email import check_new_participant_email

    seen = {}
    for row in rows:
        errors = []
        warnings = []
        email = row["email"]
        if not email:
            errors.append("missing_email")
        else:
            try:
                validate_email(email)
            except ValidationError:
                errors.append("invalid_email")
            else:
                if check_new_participant_email(email).is_disposable:
                    warnings.append("disposable_email")
        if not row["first_name"]:
            errors.append("missing_first_name")
        if not row["last_name"]:
            errors.append("missing_last_name")
        if not row["type"]:
            errors.append("invalid_type" if row["type_raw"] else "missing_type")
        if email and "invalid_email" not in errors:
            if email in seen:
                errors.append(f"duplicate_email_in_file:row_{seen[email]}")
            else:
                seen[email] = row["row"]
        row["errors"] = errors
        row["warnings"] = warnings
    return rows


# ---------------------------------------------------------------------------
# Event / track checks
# ---------------------------------------------------------------------------

def check_event(event):
    if event.status == "cancelled":
        raise ImportFileError("This event is cancelled.", code="event_cancelled")


def resolve_track(event, track_id):
    """Application events need an open track with an active pricing tier."""
    if event.registration_type != "apply":
        return None
    if not track_id:
        raise ImportFileError("Select an application track.", code="track_required")
    track = event.application_tracks.filter(pk=track_id).first()
    if not track:
        raise ImportFileError("Track not found for this event.", code="track_not_found")
    if not track.is_active or track.status != "open":
        raise ImportFileError("The selected track is not open.", code="track_not_open")
    if not track.pricing_tiers.filter(is_active=True).exists():
        raise ImportFileError("The selected track has no active pricing tier.", code="track_no_tier")
    return track


def track_summary(track):
    if not track:
        return None
    return {
        "id": track.id,
        "label": track.label,
        "roles": [role for role in (track.role_mappings_on_acceptance or []) if role],
    }


# ---------------------------------------------------------------------------
# Planning
# ---------------------------------------------------------------------------

def _find_user(email):
    matches = list(User.objects.filter(email__iexact=email).order_by("id")[:2])
    if len(matches) > 1:
        return None, "multiple_accounts"
    return (matches[0] if matches else None), None


def plan_row(event, row, track):
    """
    Decide what to do with one validated row. Read-only.

    Returns a dict: action, account ("create"/"existing"/None), application
    ("new"/"existing"/None), reason, warnings, user (User or None).
    """
    plan = {
        "action": None, "account": None, "application": None, "reason": "",
        "warnings": list(row.get("warnings", [])), "user": None,
    }

    if row.get("errors"):
        plan.update(action=ACTION_ERROR, reason=row["errors"][0])
        return plan

    user, problem = _find_user(row["email"])
    if problem:
        plan.update(action=ACTION_ERROR, reason=problem)
        return plan
    plan["user"] = user
    plan["account"] = "existing" if user else "create"

    registration = (
        EventRegistration.objects.filter(event=event, user=user).first() if user else None
    )
    is_registered = bool(registration and registration.status == "registered")

    if event.registration_type != "apply":
        if is_registered:
            plan.update(action=ACTION_SKIP, reason="already_registered")
        else:
            plan["action"] = ACTION_REGISTER
            if registration:
                plan["warnings"].append(f"reactivates_{registration.status}_registration")
        return plan

    application = EventApplication.objects.filter(event=event, email__iexact=row["email"]).first()
    if application is None and user:
        application = EventApplication.objects.filter(event=event, user=user).first()
    if application and application.user_id and user and application.user_id != user.id:
        plan.update(action=ACTION_ERROR, reason="application_linked_to_other_account")
        return plan

    track_application = None
    if application:
        track_application = application.track_applications.filter(track=track).order_by("-id").first()

    if track_application and track_application.status == EventApplicationTrackApplication.STATUS_ACCEPTED and is_registered:
        plan.update(action=ACTION_SKIP, reason="already_accepted")
        return plan

    plan["action"] = ACTION_ACCEPT
    plan["application"] = "existing" if track_application else "new"
    if track_application and track_application.status in (
        EventApplicationTrackApplication.STATUS_DECLINED,
        EventApplicationTrackApplication.STATUS_CANCELLED,
    ):
        plan["warnings"].append(f"was_{track_application.status}")
    return plan


def _serialize_plan_row(row, plan):
    return {
        "row": row["row"],
        "email": row["email"],
        "first_name": row["first_name"],
        "last_name": row["last_name"],
        "type": row["type"],
        "type_label": row.get("type_label") or row.get("type_raw", ""),
        "action": plan["action"],
        "account": plan["account"],
        "application": plan["application"],
        "reason": plan["reason"],
        "warnings": plan["warnings"],
    }


def _summarize(serialized_rows):
    summary = {
        "total": len(serialized_rows),
        "create_account": 0,
        "existing_account": 0,
        "register": 0,
        "accept_application": 0,
        "skip": 0,
        "error": 0,
    }
    for item in serialized_rows:
        action = item["action"]
        if action == ACTION_REGISTER:
            summary["register"] += 1
        elif action == ACTION_ACCEPT:
            summary["accept_application"] += 1
        elif action == ACTION_SKIP:
            summary["skip"] += 1
        else:
            summary["error"] += 1
        if action in (ACTION_REGISTER, ACTION_ACCEPT):
            if item["account"] == "create":
                summary["create_account"] += 1
            else:
                summary["existing_account"] += 1
    return summary


def capacity_warning(event, new_registrations):
    if event.max_participants is None or not new_registrations:
        return None
    current = event.registrations.filter(status="registered").count()
    if current + new_registrations <= event.max_participants:
        return None
    return {
        "max_participants": event.max_participants,
        "current": current,
        "after_import": current + new_registrations,
    }


def build_preview(event, upload, track_id, admin):
    """Parse, validate and plan the file; cache the plan. Returns the preview payload."""
    check_event(event)
    track = resolve_track(event, track_id)
    rows = validate_rows(resolve_event_types(parse_file(upload), event))

    serialized = [_serialize_plan_row(row, plan_row(event, row, track)) for row in rows]
    summary = _summarize(serialized)

    plan_id = uuid.uuid4().hex
    cache.set(
        PLAN_CACHE_KEY.format(plan_id),
        {
            "event_id": event.id,
            "admin_id": admin.id,
            "track_id": track.id if track else None,
            "rows": rows,
        },
        PLAN_TTL,
    )
    return {
        "plan_id": plan_id,
        "expires_in": PLAN_TTL,
        "event": {
            "id": event.id,
            "title": event.title,
            "registration_type": event.registration_type,
            "track": track_summary(track),
        },
        "summary": summary,
        "capacity_warning": capacity_warning(event, summary["register"] + summary["accept_application"]),
        "type_options": type_options(event),
        "rows": serialized,
    }


# ---------------------------------------------------------------------------
# Apply
# ---------------------------------------------------------------------------

class RowError(Exception):
    """A row-level failure with a reason code for the report."""


def _cognito_client():
    from django.conf import settings
    import boto3

    region = getattr(settings, "COGNITO_REGION", "") or ""
    pool_id = getattr(settings, "COGNITO_USER_POOL_ID", "") or ""
    if not region or not pool_id:
        raise RowError("cognito_not_configured")
    return boto3.client("cognito-idp", region_name=region), pool_id


def _create_cognito_login(row):
    """
    Create the Cognito user (email verified, no Cognito email, permanent password).

    Unlike create_cognito_user(), an existing Cognito username is never reused:
    that helper resets the existing user's password, which here could take over
    someone's self-signup. Uses only AdminCreateUser/AdminSetUserPassword, the
    same calls as admin Create User. Returns (username, temp_password).
    """
    from botocore.exceptions import ClientError
    from users.email_utils import generate_temporary_password

    client, pool_id = _cognito_client()
    base = "".join(c for c in row["email"].split("@")[0] if c.isalnum() or c in "._-") or "user"
    temp_password = generate_temporary_password()

    for attempt in range(50):
        username = base if attempt == 0 else f"{base}{attempt}"
        if User.objects.filter(username__iexact=username).exists():
            continue
        try:
            client.admin_create_user(
                UserPoolId=pool_id,
                Username=username,
                TemporaryPassword=temp_password,
                UserAttributes=[
                    {"Name": "email", "Value": row["email"]},
                    {"Name": "email_verified", "Value": "true"},
                    {"Name": "given_name", "Value": row["first_name"]},
                    {"Name": "family_name", "Value": row["last_name"]},
                ],
                MessageAction="SUPPRESS",
            )
        except ClientError as exc:
            code = exc.response.get("Error", {}).get("Code", "")
            if code == "UsernameExistsException":
                continue
            if code == "AliasExistsException":
                # A verified Cognito login already owns this email but has no DB account.
                raise RowError("email_exists_in_login_system")
            logger.error("Cognito create failed for import row %s: %s", row["row"], exc)
            raise RowError("cognito_create_failed")

        try:
            client.admin_set_user_password(
                UserPoolId=pool_id, Username=username, Password=temp_password, Permanent=True
            )
        except ClientError as exc:
            logger.error("Cognito set password failed for import row %s: %s", row["row"], exc)
            _delete_cognito_login(username)
            raise RowError("cognito_create_failed")
        return username, temp_password

    raise RowError("username_unavailable")


def _delete_cognito_login(username):
    from users.email_utils import delete_cognito_user

    try:
        delete_cognito_user(username)
    except Exception:
        logger.exception("Could not remove Cognito user %s after a failed import row", username)


def _create_account(row):
    """Create the Cognito user first, then the Django user. Returns (user, temp_password)."""
    username, temp_password = _create_cognito_login(row)

    try:
        with transaction.atomic():
            user = User.objects.create_user(
                username=username,
                email=row["email"],
                password=temp_password,
                first_name=row["first_name"],
                last_name=row["last_name"],
            )
            _fill_profile(user, row, new_account=True)
    except Exception:
        # Never leave a Cognito login without a matching DB account
        _delete_cognito_login(username)
        raise
    return user, temp_password


def _fill_profile(user, row, new_account=False):
    """New accounts get the file's details; existing profiles only gain empty fields."""
    from users.models import UserProfile

    profile, _ = UserProfile.objects.get_or_create(user=user)
    updates = {
        "full_name": f"{row['first_name']} {row['last_name']}".strip(),
        "job_title": row["job_title"],
        "company": row["company"],
        "location": row["country"],
    }
    changed = []
    for field, value in updates.items():
        if value and (new_account or not (getattr(profile, field) or "").strip()):
            setattr(profile, field, value[:255])
            changed.append(field)
    if changed:
        profile.save(update_fields=changed)


def _add_type_label(event, registration, row):
    from events.models import EventBadgeLabel

    name = row["type_label"]
    label = EventBadgeLabel.objects.filter(event=event, name__iexact=name).first()
    if label is None:
        color = TYPE_LABEL_COLORS.get(row["type"], EVENT_LABEL_COLOR)
        label = EventBadgeLabel.objects.create(event=event, name=name, color=color)
    elif not label.is_active:
        label.is_active = True
        label.deactivated_at = None
        label.deactivated_by = None
        label.deactivation_reason = ""
        label.save(update_fields=["is_active", "deactivated_at", "deactivated_by", "deactivation_reason"])
    registration.badge_labels.add(label)


def _mark_registration_paid(registration, admin, payment_reference):
    registration.marked_paid_by = admin
    registration.marked_paid_at = timezone.now()
    fields = ["marked_paid_by", "marked_paid_at"]
    if payment_reference:
        registration.payment_reference = payment_reference
        fields.append("payment_reference")
    registration.save(update_fields=fields)


def _register_open_event(event, user, admin):
    from events.views import _grant_invited_event_access

    return _grant_invited_event_access(event, user, invited_by=admin)["registration"]


def _accept_application(event, user, row, track, admin):
    """Reuse or create the application on the track, then accept it (no acceptance email)."""
    from events.services.application_decisions import accept_track_application

    now = timezone.now()
    note = f"Imported by {admin.email or admin.username} on {now:%Y-%m-%d} (type: {row['type_label']})"

    application = EventApplication.objects.filter(event=event, email__iexact=row["email"]).first()
    if application is None:
        application = EventApplication.objects.filter(event=event, user=user).first()
    if application is None:
        application = EventApplication.objects.create(
            event=event,
            user=user,
            first_name=row["first_name"],
            last_name=row["last_name"],
            email=row["email"],
            job_title=row["job_title"],
            company_name=row["company"],
            location=row["country"],
            submission_mode=EventApplication.SUBMISSION_MODE_SELF,
            application_track=track,
            status="pending",
        )

    application.user = user
    application.status = "approved"
    application.reviewed_by = admin
    application.reviewed_at = now
    application.is_preapproved = True
    application.preapproved_at = now
    application.comments = f"{application.comments}\n{note}".strip() if application.comments else note
    application.save(update_fields=[
        "user", "status", "reviewed_by", "reviewed_at", "is_preapproved", "preapproved_at", "comments",
    ])

    track_application = application.track_applications.filter(track=track).order_by("-id").first()
    if track_application is None:
        track_application = EventApplicationTrackApplication.objects.create(
            application=application,
            track=track,
            submission_mode=application.submission_mode or EventApplication.SUBMISSION_MODE_SELF,
        )

    accept_track_application(track_application, admin, send_email=False)
    return EventRegistration.objects.get(event=event, user=user)


def _confirm_paid_origins(registration, track, admin, payment_reference):
    """
    Paid Attendee on a paid tier: confirm this track's payment_pending origins
    (paid offline). Same result as mark_origin_paid, without its payment email.
    """
    from events.services.attendee_directory import _recalculate_registration_status

    pending = registration.origins.filter(track=track, status="active", origin_status="payment_pending")
    if not pending.exists():
        return
    update = {"origin_status": "confirmed", "marked_paid_by": admin, "marked_paid_at": timezone.now()}
    if payment_reference:
        update["payment_reference"] = payment_reference
    pending.update(**update)
    # Moving to "confirmed" triggers the post-acceptance forms via the
    # EventRegistration post_save signal.
    _recalculate_registration_status(registration)


def apply_row(event, row, track, admin):
    """Re-plan and execute one row. Returns the row's result dict (never raises)."""
    from users.email_utils import send_attendee_account_email, send_user_registration_acknowledgement_email

    plan = plan_row(event, row, track)
    result = _serialize_plan_row(row, plan)
    result.update(status=plan["action"], emails={})
    if plan["action"] in (ACTION_SKIP, ACTION_ERROR):
        return result

    temp_password = None
    try:
        user = plan["user"]
        if user is None:
            user, temp_password = _create_account(row)
        else:
            _fill_profile(user, row)

        with transaction.atomic():
            if plan["action"] == ACTION_REGISTER:
                registration = _register_open_event(event, user, admin)
            else:
                registration = _accept_application(event, user, row, track, admin)
            _add_type_label(event, registration, row)
            if row["type"] == TYPE_PAID_ATTENDEE:
                if track:
                    _confirm_paid_origins(registration, track, admin, row["payment_reference"])
                _mark_registration_paid(registration, admin, row["payment_reference"])
    except Exception as exc:
        if isinstance(exc, RowError):
            reason = str(exc)
        else:
            reason = "unexpected_error"
            logger.exception("Attendee import row %s failed for event %s", row["row"], event.id)
        result.update(status=ACTION_ERROR, reason=reason)
        if temp_password:
            # The account exists; the person can still log in with the emailed password.
            result["account_created"] = True
            result["emails"]["credentials"] = _safe_send(send_attendee_account_email, user, temp_password, event)
        return result

    result["user_id"] = user.id
    result["registration_id"] = registration.id
    if temp_password:
        result["emails"]["credentials"] = _safe_send(send_attendee_account_email, user, temp_password, event)
    result["emails"]["registration_confirmation"] = _safe_send(
        send_user_registration_acknowledgement_email, user, event
    )
    return result


def _safe_send(func, *args):
    try:
        return bool(func(*args))
    except Exception:
        logger.exception("Attendee import email %s failed", getattr(func, "__name__", func))
        return False


# ---------------------------------------------------------------------------
# Jobs
# ---------------------------------------------------------------------------

def start_job(event, plan_id, admin):
    """Take (single use) the cached plan and queue the job. Returns the job_id."""
    plan_key = PLAN_CACHE_KEY.format(plan_id)
    plan = cache.get(plan_key)
    if not plan or plan["event_id"] != event.id:
        raise ImportFileError("This preview has expired. Upload the file again.", code="plan_expired")
    if plan["admin_id"] != admin.id:
        raise ImportFileError("This preview belongs to another admin.", code="plan_not_owned")
    cache.delete(plan_key)

    job_id = uuid.uuid4().hex
    cache.set(
        JOB_CACHE_KEY.format(job_id),
        {
            "job_id": job_id,
            "event_id": event.id,
            "admin_id": admin.id,
            "track_id": plan["track_id"],
            "status": "queued",
            "total": len(plan["rows"]),
            "processed": 0,
            "results": [],
            "summary": None,
            "rows": plan["rows"],
        },
        JOB_TTL,
    )

    from events.tasks import run_attendee_import_task
    transaction.on_commit(lambda: run_attendee_import_task.delay(job_id))
    return job_id


def run_job(job_id):
    """Celery entry point: apply every row of the job, saving progress to cache."""
    from events.models import Event

    key = JOB_CACHE_KEY.format(job_id)
    job = cache.get(key)
    if not job or job["status"] not in ("queued",):
        return
    event = Event.objects.get(pk=job["event_id"])
    admin = User.objects.get(pk=job["admin_id"])
    track = event.application_tracks.filter(pk=job["track_id"]).first() if job["track_id"] else None

    job["status"] = "running"
    cache.set(key, job, JOB_TTL)
    try:
        check_event(event)
        if event.registration_type == "apply" and track is None:
            raise ImportFileError("The selected track no longer exists.", code="track_not_found")
        for row in job["rows"]:
            job["results"].append(apply_row(event, row, track, admin))
            job["processed"] += 1
            cache.set(key, job, JOB_TTL)
    except ImportFileError as exc:
        job.update(status="failed", error=str(exc), error_code=exc.code)
        cache.set(key, job, JOB_TTL)
        return

    job["summary"] = _result_summary(job["results"])
    job["status"] = "done"
    cache.set(key, job, JOB_TTL)
    logger.info("Attendee import job %s for event %s done: %s", job_id, event.id, job["summary"])


def _result_summary(results):
    summary = {"total": len(results), "registered": 0, "accounts_created": 0, "skipped": 0, "errors": 0, "emails_failed": 0}
    for item in results:
        status = item["status"]
        if status in (ACTION_REGISTER, ACTION_ACCEPT):
            summary["registered"] += 1
        elif status == ACTION_SKIP:
            summary["skipped"] += 1
        else:
            summary["errors"] += 1
        if "credentials" in item["emails"]:
            summary["accounts_created"] += 1
        summary["emails_failed"] += sum(1 for sent in item["emails"].values() if not sent)
    return summary


def get_job(event, job_id):
    job = cache.get(JOB_CACHE_KEY.format(job_id))
    if not job or job["event_id"] != event.id:
        return None
    public = {k: v for k, v in job.items() if k != "rows"}
    return public


def template_csv():
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(TEMPLATE_HEADER)
    writer.writerow(TEMPLATE_EXAMPLE)
    return buffer.getvalue()
