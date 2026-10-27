import csv
import io
import os
import tempfile
from datetime import timedelta
from unittest.mock import Mock, patch

from django.contrib.auth import get_user_model
from django.core.files.storage import default_storage
from django.test import SimpleTestCase, TestCase, override_settings
from django.urls import resolve, reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from newsletter import contact_export_services as svc
from newsletter.mautic import TemporaryMauticError
from newsletter.models import MauticContactMapping, MauticIdentityAuditLog, NewsletterCategory, NewsletterSubscription
from newsletter.tests.marketing_actors import grant_marketing_access

User = get_user_model()
TMP = tempfile.mkdtemp(prefix="contact-export-tests-")
ISOLATED = {
    "CACHES": {"default": {"BACKEND": "django.core.cache.backends.locmem.LocMemCache", "LOCATION": "contact-export-tests"}},
    "STORAGES": {
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage", "OPTIONS": {"location": TMP}},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    },
}
FIELDS = {"fields": [
    {"alias": "lastname", "label": "Last Name", "object": "lead", "order": 7, "id": 3},
    {"alias": "firstname", "label": "First Name", "object": "lead", "order": 5, "id": 2},
    {"alias": "email", "label": "Email", "object": "lead", "order": 13, "id": 4},
    {"alias": "points", "label": "Points", "object": "lead", "order": 20, "id": 5},
    {"alias": "linkedin", "label": "mautic.lead.field.linkedin", "object": "lead", "order": 40, "id": 6},
    {"alias": "companyemail", "label": "Company Email", "object": "company", "order": 8, "id": 7},
]}


def contact(cid, first="Ada", **extra):
    data = {
        "id": cid, "points": 5, "lastActive": "2026-10-20 10:00:00", "dateAdded": "2026-01-01 09:00:00",
        "dateIdentified": None, "owner": None, "tags": [{"tag": "vip"}, {"tag": "event"}],
        "stage": {"id": 9, "name": "Subscriber"}, "doNotContact": [],
        "fields": {"all": {"firstname": first, "lastname": "Lovelace", "email": f"c{cid}@example.test", "points": 5}},
    }
    data.update(extra)
    return data


def page(contacts, total=None):
    return {"total": len(contacts) if total is None else total, "contacts": {str(c["id"]): c for c in contacts}}


class ExportRulesTests(SimpleTestCase):
    def test_cells_never_start_a_formula_and_hold_no_xml_control_characters(self):
        self.assertEqual([svc._cell(v) for v in ("=1+1", "+44", "-3", "@x", "\tx", "ok", None, True, ["a", "", "b"], "a\x00b")],
                         ["'=1+1", "'+44", "'-3", "'@x", "'\tx", "ok", "", "Yes", "a | b", "ab"])

    def test_request_parsing(self):
        self.assertEqual(svc.parse_request({"format": "XLSX", "search": " ada ", "stage_id": "none"})["filters"],
                         {"search": "ada", "stage_id": "none"})
        self.assertEqual(svc.parse_request({"contact_ids": [3, "3", "12"]})["contact_ids"], ["3", "12"])
        for data, code in (({"format": "pdf"}, "invalid_format"), ({"contact_ids": "1,2"}, "invalid_contact_ids"),
                           ({"contact_ids": [True]}, "invalid_contact_ids"), ({"contact_ids": list(range(1, 1002))}, "too_many_selected"),
                           ({"stage_id": "x"}, "invalid_stage")):
            with self.subTest(data=data), self.assertRaises(svc.ContactExportError) as raised:
                svc.parse_request(data)
            self.assertEqual(raised.exception.code, code)

    @patch("newsletter.contact_export_services.MauticClient")
    @patch("newsletter.contact_export_services.marketing_cache.contact_fields", return_value=FIELDS)
    def test_columns_are_contact_fields_in_order_then_extras_without_duplicates(self, *_):
        labels = [c["label"] for c in svc.export_columns()]
        self.assertEqual(labels[:6], ["Mautic ID", "First Name", "Last Name", "Email", "Points", "Linkedin"])
        self.assertNotIn("Company Email", labels)
        self.assertEqual(labels.count("Points"), 1)
        self.assertEqual(labels[6:], ["Stage", "Tags", "Email Do Not Contact", "Owner", "Last Active", "Date Added",
                                      "Date Identified", "IMAA Connect Account", "Subscription Lists"])

    def test_row_values(self):
        columns = [{"key": k, "label": k} for k in ("id", "field:firstname", "stage", "tags", "dnc_email", "owner", "ecp_user", "lists")]
        row = svc.contact_row(
            contact(4, first="=cmd", doNotContact=[{"channel": "email", "reason": 1}, {"channel": "sms", "reason": 3}],
                    owner={"firstName": "Sam", "lastName": "Lee"}),
            columns, {"ecp_user": "Yes", "lists": "Weekly"},
        )
        self.assertEqual(row, ["4", "'=cmd", "Subscriber", "vip | event", "Unsubscribed", "Sam Lee", "Yes", "Weekly"])

    def test_filtered_pages_use_an_id_cursor_after_the_stage_filter(self):
        client = Mock()
        client.list_contacts.side_effect = [page([contact(i) for i in range(1, 201)], total=250), page([contact(i) for i in range(201, 251)], total=250)]
        totals = []
        pages = list(svc.iter_contacts({"scope": "filter", "filters": {"search": "ada", "stage_id": "9"}}, client, on_total=totals.append))
        self.assertEqual([len(p) for p in pages], [200, 50])
        self.assertEqual(totals, [250])
        second = client.list_contacts.call_args_list[1].kwargs
        self.assertEqual((second["where[0][col]"], second["where[0][val]"]), ("stage_id", "9"))
        self.assertEqual((second["where[1][col]"], second["where[1][expr]"], second["where[1][val]"]), ("id", "gt", 200))
        self.assertEqual((second["search"], second["orderBy"]), ("ada", "id"))

    def test_too_many_contacts_are_refused_before_any_row_is_built(self):
        client = Mock()
        client.list_contacts.return_value = page([contact(1)], total=100_001)
        with self.assertRaises(svc.ContactExportError) as raised:
            list(svc.iter_contacts({"scope": "filter", "filters": {}}, client))
        self.assertEqual(raised.exception.code, "export_too_large")

    def test_selected_contacts_are_fetched_by_id(self):
        client = Mock()
        client.list_contacts.return_value = page([contact(3), contact(12)])
        list(svc.iter_contacts({"scope": "selected", "contact_ids": ["3", "12"]}, client))
        self.assertEqual(client.list_contacts.call_args.kwargs["search"], "ids:3,12")


@override_settings(**ISOLATED)
class ExportJobTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(username="export-admin", email="ea@example.test", password="x")
        grant_marketing_access(self.user)
        for target, kwargs in (
            ("newsletter.contact_export_services.marketing_cache.contact_fields", {"return_value": FIELDS}),
            ("newsletter.contact_export_services.MauticClient", {}),
        ):
            patcher = patch(target, **kwargs)
            patcher.start()
            self.addCleanup(patcher.stop)
        self.client_mock = Mock()
        self.client_mock.list_contacts.return_value = page([contact(1), contact(2, first="Grace")])
        category = NewsletterCategory.objects.create(name="Weekly", slug="weekly-export")
        member = User.objects.create_user(username="export-member", email="c1@example.test", password="x")
        MauticContactMapping.objects.create(user=member, mautic_contact_id="1")
        NewsletterSubscription.objects.create(user=member, category=category, is_subscribed=True)

    def run_job(self, fmt="csv", **options):
        job = svc.create_job(self.user, svc.parse_request({"format": fmt, **options}))
        return svc.run_export(job["id"], client=self.client_mock)

    def test_csv_export_is_stored_encrypted_and_decrypts_for_download(self):
        job = self.run_job()
        self.assertEqual((job["status"], job["rows"], job["total"]), ("done", 2, 2))
        with default_storage.open(job["storage_name"], "rb") as handle:
            self.assertNotIn(b"c1@example.test", handle.read())
        rows = list(csv.reader(io.StringIO(svc.read_file(job).decode("utf-8-sig"))))
        self.assertEqual(rows[0][:3], ["Mautic ID", "First Name", "Last Name"])
        self.assertEqual(rows[1][rows[0].index("Subscription Lists")], "Weekly")
        self.assertEqual(rows[1][rows[0].index("IMAA Connect Account")], "Yes")
        self.assertEqual(rows[2][rows[0].index("IMAA Connect Account")], "No")
        self.assertNotIn("key", svc.public_job(job))

    def test_xlsx_export_keeps_formula_like_text_as_text(self):
        self.client_mock.list_contacts.return_value = page([contact(1, first="=HYPERLINK(\"x\")")])
        job = self.run_job("xlsx")
        sheet = load_workbook(io.BytesIO(svc.read_file(job))).active
        self.assertEqual(sheet.cell(row=2, column=2).value, "'=HYPERLINK(\"x\")")
        self.assertEqual(sheet.cell(row=2, column=2).data_type, "s")

    def test_mautic_failure_ends_in_a_failed_job(self):
        self.client_mock.list_contacts.side_effect = TemporaryMauticError("down")
        job = self.run_job()
        self.assertEqual((job["status"], job["storage_name"]), ("failed", ""))
        self.assertIn("Mautic could not be read", job["error"])

    def test_a_missing_or_tampered_file_is_reported_gone(self):
        job = self.run_job()
        default_storage.delete(job["storage_name"])
        with self.assertRaises(svc.ContactExportError) as raised:
            svc.read_file(job)
        self.assertEqual(raised.exception.status_code, 410)

    def test_purge_deletes_only_files_past_retention(self):
        old, fresh = self.run_job(), self.run_job()
        past = (timezone.now() - timedelta(hours=25)).timestamp()
        os.utime(os.path.join(TMP, old["storage_name"]), (past, past))
        self.assertEqual(svc.purge_expired_files(), 1)
        self.assertFalse(default_storage.exists(old["storage_name"]))
        self.assertTrue(default_storage.exists(fresh["storage_name"]))


@override_settings(**ISOLATED)
class ExportApiTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username="export-api", email="eapi@example.test", password="x")
        grant_marketing_access(self.admin)
        self.api = APIClient()
        self.api.force_authenticate(self.admin)
        self.list_url = reverse("newsletter-admin-contact-export-list")
        patcher = patch("newsletter.tasks.run_contact_export.delay")
        self.delay = patcher.start()
        self.addCleanup(patcher.stop)

    def test_exports_route_is_not_a_contact_id(self):
        self.assertEqual(resolve("/api/newsletter/admin/contacts/exports/").url_name, "newsletter-admin-contact-export-list")

    def test_start_queues_the_job_and_audits_it(self):
        r = self.api.post(self.list_url, {"format": "xlsx", "stage_id": "9"}, format="json")
        self.assertEqual(r.status_code, 202)
        self.delay.assert_called_once_with(r.json()["id"])
        self.assertEqual((r.json()["status"], r.json()["scope"]), ("queued", "filter"))
        entry = MauticIdentityAuditLog.objects.get(action="contact.export.create")
        self.assertEqual(entry.detail, "xlsx, filter, stage 9")
        self.assertEqual(r["Cache-Control"], "no-store")

    def test_broker_outage_is_reported_and_leaves_no_job(self):
        self.delay.side_effect = ConnectionError("broker down")
        r = self.api.post(self.list_url, {"format": "csv"}, format="json")
        self.assertEqual((r.status_code, r.json()["code"]), (503, "export_not_queued"))
        self.assertEqual(self.api.get(self.list_url).json()["results"], [])

    def test_exports_are_private_to_the_admin_who_started_them(self):
        job_id = self.api.post(self.list_url, {"format": "csv"}, format="json").json()["id"]
        other = User.objects.create_superuser(username="export-other", email="eo@example.test", password="x")
        grant_marketing_access(other)
        api = APIClient()
        api.force_authenticate(other)
        for method, name in (("get", "newsletter-admin-contact-export-detail"), ("get", "newsletter-admin-contact-export-download"),
                             ("delete", "newsletter-admin-contact-export-detail")):
            with self.subTest(name=name, method=method):
                self.assertEqual(getattr(api, method)(reverse(name, args=[job_id])).status_code, 404)
        self.assertEqual(api.get(self.list_url).json()["results"], [])

    def test_download_before_the_file_is_ready_is_refused(self):
        job_id = self.api.post(self.list_url, {"format": "csv"}, format="json").json()["id"]
        r = self.api.get(reverse("newsletter-admin-contact-export-download", args=[job_id]))
        self.assertEqual((r.status_code, r.json()["code"]), (409, "export_not_ready"))

    def test_non_marketing_users_are_denied(self):
        normal = APIClient()
        normal.force_authenticate(User.objects.create_user(username="export-normal", email="en@example.test", password="x"))
        self.assertIn(normal.post(self.list_url, {"format": "csv"}, format="json").status_code, (401, 403))
        self.delay.assert_not_called()

    def test_the_export_task_has_its_own_long_time_limit(self):
        from newsletter.tasks import run_contact_export

        self.assertEqual((run_contact_export.time_limit, run_contact_export.soft_time_limit), (1800, 1740))


# Notifications and the stale-job watchdog ---------------------------------------

from friends.models import Notification  # noqa: E402


@override_settings(**ISOLATED)
class ExportNotificationTests(ExportJobTests):
    def test_a_finished_export_notifies_its_owner(self):
        job = self.run_job("xlsx")
        note = Notification.objects.get(recipient=self.user, data__type="marketing_contact_export")
        self.assertEqual((note.kind, note.title, note.is_read), ("system", "Contact export ready", False))
        self.assertEqual(note.description, "2 contacts exported as XLSX. Download it within 24 hours.")
        self.assertEqual((note.data["job_id"], note.data["status"], note.data["filename"]), (job["id"], "done", job["filename"]))

    def test_a_failed_export_notifies_with_the_reason(self):
        self.client_mock.list_contacts.side_effect = TemporaryMauticError("down")
        self.run_job()
        note = Notification.objects.get(recipient=self.user, data__type="marketing_contact_export")
        self.assertEqual(note.title, "Contact export failed")
        self.assertIn("Mautic could not be read", note.description)

    def test_a_job_no_worker_picked_up_fails_and_notifies_once(self):
        job = svc.create_job(self.user, svc.parse_request({"format": "csv"}))
        self.assertEqual(svc.expire_if_stale(job)["status"], "queued")
        later = timezone.now() + timedelta(minutes=4)
        stale = svc.expire_if_stale(svc.get_job(job["id"]), now=later)
        self.assertEqual(stale["status"], "failed")
        self.assertIn("background workers may not be running", stale["error"])
        svc.expire_if_stale(svc.get_job(job["id"]), now=later)
        self.assertEqual(Notification.objects.filter(recipient=self.user, data__job_id=job["id"]).count(), 1)

    def test_a_running_job_past_the_task_limit_is_reported_interrupted(self):
        job = svc.create_job(self.user, svc.parse_request({"format": "csv"}))
        svc._update_job(job["id"], status="running")
        self.assertEqual(svc.expire_if_stale(svc.get_job(job["id"]), now=timezone.now() + timedelta(minutes=10))["status"], "running")
        stale = svc.expire_if_stale(svc.get_job(job["id"]), now=timezone.now() + timedelta(minutes=36))
        self.assertIn("interrupted", stale["error"])


@override_settings(**ISOLATED)
class MarketingNotificationApiTests(TestCase):
    def setUp(self):
        self.admin = User.objects.create_superuser(username="notify-admin", email="na@example.test", password="x")
        grant_marketing_access(self.admin)
        self.api = APIClient()
        self.api.force_authenticate(self.admin)
        self.mine = Notification.objects.create(recipient=self.admin, kind="system", title="Contact export ready",
                                                data={"type": "marketing_contact_export", "job_id": "j1", "status": "done", "filename": "c.csv"})
        Notification.objects.create(recipient=self.admin, kind="system", title="Post approved", data={"type": "moderation"})
        other = User.objects.create_user(username="notify-other", email="no@example.test", password="x")
        Notification.objects.create(recipient=other, kind="system", title="Theirs", data={"type": "marketing_contact_export"})
        self.url = reverse("newsletter-admin-marketing-notifications")

    def test_only_my_marketing_notifications_are_listed(self):
        body = self.api.get(self.url).json()
        self.assertEqual(body["unread"], 1)
        self.assertEqual([n["title"] for n in body["results"]], ["Contact export ready"])
        self.assertEqual(body["results"][0]["data"], {"job_id": "j1", "status": "done", "filename": "c.csv"})

    def test_mark_read_by_id_or_all_and_validate(self):
        read_url = reverse("newsletter-admin-marketing-notifications-read")
        self.assertEqual(self.api.post(read_url, {"ids": ["x"]}, format="json").status_code, 400)
        self.assertEqual(self.api.post(read_url, {"ids": [self.mine.id]}, format="json").json(), {"updated": 1, "unread": 0})
        self.mine.refresh_from_db()
        self.assertTrue(self.mine.is_read)
        Notification.objects.filter(pk=self.mine.pk).update(is_read=False)
        self.assertEqual(self.api.post(read_url, {"all": True}, format="json").json()["updated"], 1)
        # Someone else's or non-marketing notifications are never touched.
        self.assertEqual(Notification.objects.filter(is_read=True).count(), 1)

    def test_reading_notifications_does_not_invalidate_the_marketing_cache(self):
        with patch("newsletter.middleware.marketing_cache.invalidate") as invalidate:
            self.api.post(reverse("newsletter-admin-marketing-notifications-read"), {"all": True}, format="json")
        invalidate.assert_not_called()

    def test_non_marketing_users_are_denied(self):
        normal = APIClient()
        normal.force_authenticate(User.objects.create_user(username="notify-normal", email="nn@example.test", password="x"))
        self.assertIn(normal.get(self.url).status_code, (401, 403))
